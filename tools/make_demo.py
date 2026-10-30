"""生成 20 份 / 1000 页固定课件，并独立记录预期命中列表。"""
import argparse
import json
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

TOPICS = ("操作系统", "网络基础", "系统管理", "云计算", "开发环境")


def generate_demo(root: Path):
    root.mkdir(parents=True, exist_ok=True)
    expected = []
    for file_index in range(20):
        topic = TOPICS[file_index % len(TOPICS)]
        relative = Path("基础课件" if file_index < 10 else "进阶课件") / f"{file_index + 1:02d}_{topic}.pptx"
        path = root / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        presentation = Presentation()
        for page in range(1, 51):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            title = slide.shapes.add_textbox(Inches(.7), Inches(.6), Inches(8), Inches(.7))
            title.text = f"{topic} · 第 {page} 讲"
            title.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
            text = f"课程知识点 {page}：基础概念、实践步骤与复习思考。"
            if page % 7 == 0:
                text = f"Linux 内核与进程管理：本节学习 linux kernel 和网络配置。课程 {file_index + 1}。"
                expected.append([str(relative), page, "正文", text])
            box = slide.shapes.add_textbox(Inches(.7), Inches(1.8), Inches(8), Inches(1.5))
            if page % 7 == 0:
                paragraph = box.text_frame.paragraphs[0]
                paragraph.add_run().text = text[:2]
                run = paragraph.add_run()
                run.text = text[2:]
                run.font.bold = True
            else:
                box.text = text
            if page % 11 == 0:
                notes = f"讲师提示：使用 linux 演示环境完成第 {page} 页练习。"
                slide.notes_slide.notes_text_frame.text = notes
                expected.append([str(relative), page, "备注", notes])
            if page % 13 == 0:
                table = slide.shapes.add_table(2, 2, Inches(.7), Inches(3.4), Inches(7), Inches(1)).table
                table.cell(0, 0).text = "环境"
                table.cell(0, 1).text = "命令"
                table.cell(1, 0).text = "Linux 网络配置"
                table.cell(1, 1).text = "ip address"
                expected.append([str(relative), page, "表格", "Linux 网络配置"])
            if page == 49:
                slide.element.set("show", "0")
        presentation.save(path)
    manifest = {"files": 20, "slides": 1000, "keyword": "linux", "expected": expected}
    (root / "expected-linux.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return manifest


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=Path("work/demo"))
    arguments = parser.parse_args()
    manifest = generate_demo(arguments.output)
    print(json.dumps({"directory": str(arguments.output.resolve()), "files": manifest["files"],
                      "slides": manifest["slides"], "expected_rows": len(manifest["expected"])}, ensure_ascii=False))
