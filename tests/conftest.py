from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches


def make_deck(path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[6])
    frame = first.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame
    paragraph = frame.paragraphs[0]
    paragraph.add_run().text = "Li"
    run = paragraph.add_run()
    run.text = "nux 和 linux 内核"
    run.font.bold = True
    table = first.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(5), Inches(1)).table
    table.cell(0, 0).text = "网络配置 Linux"
    table.cell(0, 1).text = "完整短语 linux kernel"
    table.cell(1, 0).merge(table.cell(1, 1))
    table.cell(1, 0).text = "合并单元格 linux"
    group = first.shapes.add_group_shape()
    nested = group.shapes.add_group_shape()
    nested.shapes.add_textbox(Inches(1), Inches(4), Inches(4), Inches(1)).text = "组合形状 linux"
    first.notes_slide.notes_text_frame.text = "讲者备注 linux 网络配置"
    first.element.set("show", "0")
    second = presentation.slides.add_slide(presentation.slide_layouts[6])
    second.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = "第二页：普通文字 a.b <tag> & 中文"
    third = presentation.slides.add_slide(presentation.slide_layouts[6])
    third.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = "第三页 Linux"
    original_first_id = first.slide_id
    # 测试故意制造“内部 slide1.xml 在第 3 页”的文件，捕捉错误页码实现。
    slide_ids = presentation.slides._sldIdLst
    node = slide_ids[0]
    slide_ids.remove(node)
    slide_ids.append(node)
    presentation.save(path)
    return original_first_id


@pytest.fixture
def deck(tmp_path):
    path = tmp_path / "中文 路径" / "操作系统.pptx"
    slide_id = make_deck(path)
    return path, slide_id
