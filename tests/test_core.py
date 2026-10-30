import csv
import json
from dataclasses import replace
from threading import Event

import pytest
from pptx import Presentation

from ppt_search.cache import TextCache
from ppt_search.extract import Cancelled, extract_document
from ppt_search.models import Document, SearchOptions
from ppt_search.presentation import export_csv, highlight
from ppt_search.search import discover_files, search_document


def test_real_order_nested_groups_tables_notes_runs_and_hidden(deck):
    path, original_first_id = deck
    document = extract_document(path)
    hits = search_document(document, SearchOptions("linux"))
    assert document.slides[2].slide_id == original_first_id
    assert document.slides[2].number == 3 and document.slides[2].hidden
    assert len(hits) == 7
    assert sum(len(hit.spans) for hit in hits) == 8
    assert {hit.block.source for hit in hits} == {"正文", "表格", "备注"}
    assert next(hit for hit in hits if "组合形状" in hit.block.text).slide_number == 3
    assert next(hit for hit in hits if "内核" in hit.block.text).block.text == "Linux 和 linux 内核"
    assert len([hit for hit in hits if "合并单元格" in hit.block.text]) == 1
    assert Document.from_dict(json.loads(json.dumps(document.to_dict()))).to_dict() == document.to_dict()


def test_literal_phrase_chinese_case_and_notes(deck):
    document = extract_document(deck[0])
    assert len(search_document(document, SearchOptions("linux kernel"))) == 1
    assert len(search_document(document, SearchOptions("网络配置"))) == 2
    assert len(search_document(document, SearchOptions("linux", case_sensitive=True))) == 5
    assert all(hit.block.source != "备注" for hit in search_document(document, SearchOptions("linux", include_notes=False)))
    assert len(search_document(document, SearchOptions("a.b"))) == 1
    assert search_document(document, SearchOptions("a*b")) == []
    with pytest.raises(ValueError):
        search_document(document, SearchOptions(" "))


def test_cache_invalidation_corruption_clear_and_prune(deck, tmp_path):
    path, _ = deck
    with TextCache(tmp_path / "cache.sqlite3") as cache:
        assert cache.get(path) is None
        document = extract_document(path)
        cache.put(document)
        assert cache.get(path).to_dict() == document.to_dict()
        cache.connection.execute("UPDATE documents SET content='invalid json'")
        cache.connection.commit()
        assert cache.get(path) is None
        cache.put(document)
        presentation = Presentation(path)
        presentation.slides[0].shapes[0].text = "编辑后的文字"
        presentation.save(path)
        assert cache.get(path) is None
        cache.put(extract_document(path))
        cache.clear()
        assert cache.get(path) is None
        cache.put(extract_document(path))
        path.unlink()
        cache.prune_missing()
        assert cache.connection.execute("SELECT count(*) FROM documents").fetchone()[0] == 0


def test_discovery_subdirectories_same_names_temp_and_cancel(deck, tmp_path):
    path, _ = deck
    copy = tmp_path / "操作系统.pptx"
    copy.write_bytes(path.read_bytes())
    (tmp_path / "~$锁定.pptx").write_bytes(b"not a presentation")
    assert discover_files(tmp_path, False) == [copy]
    assert len(discover_files(tmp_path, True)) == 2
    assert discover_files(tmp_path, True, excluded=(path.parent,)) == [copy]
    cancelled = Event()
    cancelled.set()
    with pytest.raises(Cancelled):
        discover_files(tmp_path, True, cancelled)
    with pytest.raises(Cancelled):
        extract_document(path, cancelled)


def test_html_escape_offsets_and_export(deck, tmp_path):
    assert highlight("<Linux> & linux", SearchOptions("linux")) == ('&lt;<span style="background-color:#fef08a;color:#713f12;">Linux</span>'
        '&gt; &amp; <span style="background-color:#fef08a;color:#713f12;">linux</span>')
    hit = search_document(extract_document(deck[0]), SearchOptions("linux"))[0]
    formula_hit = replace(hit, block=replace(hit.block, text="=linux"))
    target = tmp_path / "结果.csv"
    export_csv(target, [hit, formula_hit])
    assert target.read_bytes().startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(target.open(encoding="utf-8-sig", newline="")))
    assert rows[0][0] == "文件名" and rows[1][2] == str(hit.slide_number)
    assert rows[2][7] == "'=linux"
