"""只读取 PPTX；不保存、转换或修改源文件。"""
from pathlib import Path
from threading import Event

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .models import Document, SlideText, TextBlock


class Cancelled(Exception):
    """后台任务在文件或幻灯片边界取消。"""


def check_cancel(cancel: Event | None) -> None:
    if cancel is not None and cancel.is_set():
        raise Cancelled()


def paragraph_blocks(text_frame, source: str, location: str):
    # paragraph.text 会连接格式 run，因此字体或颜色切换不会拆断关键词。
    for number, paragraph in enumerate(text_frame.paragraphs, 1):
        text = paragraph.text
        if text.strip():
            yield TextBlock(text, source, f"{location} · 段落 {number}")


def shape_blocks(shapes, cancel: Event | None, parent: str = ""):
    for shape in shapes:
        check_cancel(cancel)
        location = f"{parent} / {shape.name}" if parent else shape.name
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from shape_blocks(shape.shapes, cancel, location)
            continue
        if shape.has_text_frame:
            yield from paragraph_blocks(shape.text_frame, "正文", location)
        if not shape.has_table:
            continue
        for row_number, row in enumerate(shape.table.rows, 1):
            for column_number, cell in enumerate(row.cells, 1):
                # 合并区域的被覆盖单元格不再重复提取。
                if not cell.is_spanned:
                    cell_location = f"{location} · 第 {row_number} 行 {column_number} 列"
                    yield from paragraph_blocks(cell.text_frame, "表格", cell_location)


def extract_document(path: Path, cancel: Event | None = None) -> Document:
    path = path.resolve()
    check_cancel(cancel)
    before = path.stat()
    presentation = Presentation(str(path))
    slides = []
    # 必须遍历演示文稿的顺序表，不能把 slide_id 或 slideN.xml 的 N 当页码。
    for number, slide in enumerate(presentation.slides, 1):
        check_cancel(cancel)
        title_shape = slide.shapes.title
        blocks = list(shape_blocks(slide.shapes, cancel))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None:
                blocks.extend(paragraph_blocks(notes, "备注", "讲者备注"))
        fallback_title = next((block.text.splitlines()[0][:120] for block in blocks if block.source == "正文"), "无标题")
        slides.append(SlideText(number, slide.slide_id,
                                title_shape.text if title_shape else fallback_title,
                                slide.element.get("show") in {"0", "false"}, blocks))
    after = path.stat()
    if (before.st_size, before.st_mtime_ns) != (after.st_size, after.st_mtime_ns):
        raise OSError("文件在读取过程中发生变化，请保存后重试")
    return Document(str(path), after.st_size, after.st_mtime_ns, slides)
